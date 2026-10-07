# -*- coding: utf-8 -*-
"""
Генератор 3-страничного РАБОЧЕГО ДНЕВНИКА ЭКСПЕРИМЕНТАТОРА для Алисы Ковалевой
с учетом всех методических и биологических рекомендаций.
СТРОГО РОВНО 3 СТРАНИЦЫ А4:
  Стр 1: Шапка + Регламент полива/наблюдений + Таблица Фазы 0 (Д0–Д9) с дифференцированным увлажнением.
  Стр 2: Таблица стресс-опыта Дни 10–13 (вечерний замер 18:00, 20 строк).
  Стр 3: Таблица стресс-опыта Дни 14–17 (вечерний замер 18:00, 20 строк) + Итоговый расчет репарации K_rec и подписи.
"""

import os
import subprocess
import shutil
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = r"c:\Users\Администратор\Documents\Coglet\plant-stress-ndvi"
DOCS_DIR = os.path.join(BASE_DIR, "docs")
USER_DOCS = r"C:\Users\Администратор\Documents"
CHROME_PATH = r"/usr/bin/chromium"

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=20, bottom=20, left=40, right=40):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def set_table_borders(table, color="94a3b8", sz="4", val="single"):
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def build_docx_3pages(output_path):
    doc = docx.Document()
    
    # Альбомная ориентация А4 с полями 8-10 мм
    section = doc.sections[0]
    section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(8)
    font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
    style_normal.paragraph_format.line_spacing = 1.05
    style_normal.paragraph_format.space_after = Pt(1)
    
    # =========================================================================
    # СТРАНИЦА 1
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    r = p_title.add_run("РАБОЧИЙ ДНЕВНИК ИССЛЕДОВАТЕЛЯ — КОМПЛЕКС РАННЕЙ ИНДИКАЦИИ СТРЕССА РАСТЕНИЙ\n")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x00, 0x82, 0x76)
    
    r_sub = p_title.add_run("Исследователь: Ковалева Алиса (10 кл., ГБОУ СОШ №282 СПб) | Руководитель: Ковалев И. В. | Посев: 29.09.2026 | Старт стресс-опыта: 09.10.2026")
    r_sub.font.size = Pt(8.0)
    r_sub.font.italic = True
    
    # Инструкции в 2 колонки
    t_guide = doc.add_table(rows=1, cols=2)
    t_guide.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_guide, color="0f766e", sz="6")
    c_left, c_right = t_guide.rows[0].cells[0], t_guide.rows[0].cells[1]
    c_left.width = Inches(5.4)
    c_right.width = Inches(5.4)
    set_cell_background(c_left, "f0fdf4")
    set_cell_background(c_right, "fef2f2")
    set_cell_margins(c_left, top=30, bottom=30, left=50, right=50)
    set_cell_margins(c_right, top=30, bottom=30, left=50, right=50)
    
    p_gl = c_left.paragraphs[0]
    p_gl.add_run("🌱 ФАЗА 0: ПРОРАЩИВАНИЕ (29.09 — 08.10) — ЧТО ДЕЛАТЬ И СМОТРЕТЬ:\n").bold = True
    p_gl.paragraph_format.line_spacing = 1.05
    p_gl.add_run(
        "• 29.09 (Д0): Посев семян с корешками 1–2 см. Полив отстоянной водой (15 мл). Кассета №2 СРАЗУ в отдельный лоток!\n"
        "• 30.09 (Д1): Корни без листьев не транспирируют! Не заливать болото: только увлажнить пульверизатором (5–10 мл).\n"
        "• 01.10 (Д2): СМОТРЕТЬ ПЕТЕЛЬКИ! При появлении зеленых всходов — СРАЗУ СНЯТЬ ПЛЕНКУ, чтобы не загнили!\n"
        "• 02.10 (Д3): Раскрытие семядолей. Включить подсветку (14 ч/день). Полив по 20 мл водой.\n"
        "• 03.10 (Д4): Посылка! Разворачивание 1-го настоящего листа. Монтаж светодиодов и АЦП.\n"
        "• 04–05.10 (Д5–Д6): Полив по 20 мл. На Стенде №0 настраиваем NoIR, белый эталон и OCR.\n"
        "• 06–07.10 (Д7–Д8): Полив по 20 мл. Рост 2-го настоящего листа, выравнивание проростков.\n"
        "• 08.10 (Д9): Вечер перед стартом: полив водой (20 мл), замер фона NDVI_init и массы M0."
    ).font.size = Pt(7.1)
    
    p_gr = c_right.paragraphs[0]
    p_gr.add_run("⚡ ФАЗА 1: СТРЕСС-ОПЫТ (09.10 — 16.10, замер в 18:00 на пике стресса):\n").bold = True
    p_gr.paragraph_format.line_spacing = 1.05
    p_gr.add_run(
        "• Когорта №1 (🌱 Контроль): Полив чистой водой 20 мл ЕЖЕДНЕВНО. Смотреть: тургор 100%, рост биомассы, лист прохладный (ΔT < 0).\n"
        "• Когорта №2 (🧂 Засоление 150 мМ): Полив раствором соли 150 мМ ЕЖЕДНЕВНО (20 мл, 8.76 г NaCl/л). НЕ ВЫНИМАТЬ ИЗ СВОЕГО ЛОТКА! Смотреть: почва сырая, но лист горячий (ΔT > 0).\n"
        "• Когорта №3 (🔬 Превентивная регидратация): С 09.10 НЕ ПОЛИВАТЬ! Ждать алерта (ΔT > +0.8°C, спад NDVI > 10%). При алерте — СРОЧНО РЕАНИМАЦИОННЫЙ ПОЛИВ ВОДОЙ 25 мл ДО ПОЯВЛЕНИЯ УВЯДАНИЯ!\n"
        "• Когорта №4 (👁️ Традиционный визуальный контроль): С 09.10 НЕ ПОЛИВАТЬ! Алерт игнорировать. Полив водой (25 мл) ТОЛЬКО КОГДА ЛИСТЬЯ ВИДИМО ПОВИСНУТ (угол > 30°). Смотреть: краевой некроз.\n"
        "• Когорта №5 (⚠️ Терминал): ВООБЩЕ НЕ ПОЛИВАТЬ ДО КОНЦА! Фиксация точки гибели ткани.\n"
        "*Совет по регидратации (К3 и К4): сухой торф поливать медленно шприцем под стебли, чтобы вода впиталась в ком!"
    ).font.size = Pt(7.0)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    
    # Таблица Фазы 0 (Дни 0–9)
    p_t0_title = doc.add_paragraph()
    r_t0 = p_t0_title.add_run("ТАБЛИЦА РЕАЛЬНЫХ ЗАМЕРОВ: ФАЗА 0 — ПРОРАЩИВАНИЕ И ПОДГОТОВКА (ДНИ 0–9)")
    r_t0.bold = True
    r_t0.font.size = Pt(8.0)
    r_t0.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
    
    headers_p0 = ["Дата / День", "Когорта (Кассета)", "Время", "Полив (мл / состав)", "Масса M (г)", "Всхожесть (из 9)", "Tвозд / RHвозд", "Фенонаблюдения Алисы (ростки, петельки, листья)", "Подпись"]
    t0 = doc.add_table(rows=11, cols=len(headers_p0))
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t0, color="94a3b8")
    
    for j, h in enumerate(headers_p0):
        cell = t0.rows[0].cells[j]
        set_cell_background(cell, "e2e8f0")
        set_cell_margins(cell, top=20, bottom=20, left=30, right=30)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(7.0)
        
    p0_days = [
        ("29.09 (Д0)", "К1–К5 (все)", "Вода 15–20 мл"),
        ("30.09 (Д1)", "К1–К5 (все)", "Пульверизатор 5–10мл"),
        ("01.10 (Д2)", "К1–К5 (все)", "Пульверизатор 5–10мл"),
        ("02.10 (Д3)", "К1–К5 (все)", "Вода 20 мл"),
        ("03.10 (Д4)", "К1–К5 + Ст0", "Вода 20 мл"),
        ("04.10 (Д5)", "К1–К5 + Ст0", "Вода 20 мл"),
        ("05.10 (Д6)", "К1–К5 + Ст0", "Вода 20 мл"),
        ("06.10 (Д7)", "К1–К5 (все)", "Вода 20 мл"),
        ("07.10 (Д8)", "К1–К5 (все)", "Вода 20 мл"),
        ("08.10 (Д9)", "К1–К5 (все)", "Вода 20 мл")
    ]
    
    for r_i, (d_lbl, c_lbl, pol_lbl) in enumerate(p0_days):
        row = t0.rows[r_i + 1]
        tr_pr = row._tr.get_or_add_trPr()
        tr_height = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="230" w:hRule="atLeast"/>')
        tr_pr.append(tr_height)
        
        row.cells[0].paragraphs[0].add_run(d_lbl).font.size = Pt(7.2)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[1].paragraphs[0].add_run(c_lbl).font.size = Pt(7.2)
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Предзаполненный полив
        p_pol = row.cells[3].paragraphs[0]
        p_pol.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_pol = p_pol.add_run(pol_lbl)
        r_pol.font.size = Pt(7.2)
        if "Пульверизатор" in pol_lbl:
            r_pol.font.color.rgb = RGBColor(0x02, 0x84, 0xc7)
        else:
            r_pol.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
            
        for c_cell in row.cells:
            set_cell_margins(c_cell, top=20, bottom=20, left=30, right=30)
            
    doc.add_page_break()
    
    # =========================================================================
    # СТРАНИЦА 2: ДНИ 10–13
    # =========================================================================
    p_s2_head = doc.add_paragraph()
    r = p_s2_head.add_run("СТРАНИЦА 2. ЖУРНАЛ СТРЕСС-ОПЫТА: ДНИ 10–13 (ВХОД В СТРЕСС, АЛЕРТ И ПРЕВЕНТИВНАЯ РЕГИДРАТАЦИЯ К3)")
    r.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
    
    headers_stress = [
        "День / Время", "Когорта (Кассета)", "Полив (мл/чем)", "Масса (г)", "W почвы%", "V почвы", 
        "T листа°C", "T возд°C", "ΔT (°C)", "RH%", "NDVI", "PLA см²", "Тургор / Симптомы (осмотр Алисы)", "Подпись"
    ]
    
    cohorts_template = [
        ("№1: 🌱 Контроль (Оптимум)", "Вода 20мл", "water"),
        ("№2: 🧂 Засоление (150 мМ)", "NaCl 20мл", "salt"),
        ("№3: 🔬 Превентивная регидратация", "Вода «__» мл", "device"),
        ("№4: 👁️ Традиционный визуальный контроль", "Вода «__» мл", "eyes"),
        ("№5: ⚠️ Терминал", "Полива нет", "drought")
    ]
    
    def fill_stress_table(days_list):
        t = doc.add_table(rows=1, cols=len(headers_stress))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, color="94a3b8")
        
        for j, h in enumerate(headers_stress):
            cell = t.rows[0].cells[j]
            set_cell_background(cell, "e2e8f0")
            set_cell_margins(cell, top=15, bottom=15, left=20, right=20)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h)
            r.bold = True
            r.font.size = Pt(6.8)
            
        for day_lbl in days_list:
            for c_name, c_poliv, c_type in cohorts_template:
                row = t.add_row()
                tr_pr = row._tr.get_or_add_trPr()
                tr_height = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="230" w:hRule="atLeast"/>')
                tr_pr.append(tr_height)
                
                # Дата
                row.cells[0].paragraphs[0].add_run(day_lbl).font.size = Pt(6.8)
                row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                # Когорта
                p_c = row.cells[1].paragraphs[0]
                r_c = p_c.add_run(c_name)
                r_c.font.size = Pt(7.0)
                if c_type == "water": r_c.font.color.rgb = RGBColor(0x04, 0x78, 0x57)
                elif c_type == "salt": r_c.font.color.rgb = RGBColor(0xb4, 0x53, 0x09)
                elif c_type == "device": r_c.font.color.rgb = RGBColor(0x02, 0x84, 0xc7)
                elif c_type == "eyes": r_c.font.color.rgb = RGBColor(0x6d, 0x28, 0xd9)
                elif c_type == "drought": r_c.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)
                
                # Полив (предзаполненный)
                p_p = row.cells[2].paragraphs[0]
                p_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_p = p_p.add_run(c_poliv)
                r_p.font.size = Pt(6.8)
                if c_type == "drought":
                    r_p.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c)
                elif "«__»" in c_poliv:
                    r_p.font.color.rgb = RGBColor(0x02, 0x84, 0xc7)
                else:
                    r_p.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
                    
                for c_cell in row.cells:
                    set_cell_margins(c_cell, top=15, bottom=15, left=20, right=20)
                    
    # Страница 2: Дни 10, 11, 12, 13 (с указанием вечернего времени 18:00)
    fill_stress_table(["День 10 09.10 (18:00)", "День 11 10.10 (18:00)", "День 12 11.10 (18:00)", "День 13 12.10 (18:00)"])
    
    doc.add_page_break()
    
    # =========================================================================
    # СТРАНИЦА 3: ДНИ 14–17 + ИТОГИ
    # =========================================================================
    p_s3_head = doc.add_paragraph()
    r = p_s3_head.add_run("СТРАНИЦА 3. ЖУРНАЛ СТРЕСС-ОПЫТА: ДНИ 14–17 (РЕАКТИВНЫЙ ПОЛИВ К4, РЕПАРАЦИЯ, ФИНАЛЬНЫЙ УЧЕТ)")
    r.bold = True
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
    
    fill_stress_table(["День 14 13.10 (18:00)", "День 15 14.10 (18:00)", "День 16 15.10 (18:00)", "День 17 16.10 (18:00)"])
    
    # Итоговый расчетный блок
    p_sum_title = doc.add_paragraph()
    p_sum_title.paragraph_format.space_before = Pt(3)
    p_sum_title.paragraph_format.space_after = Pt(1)
    r_sum = p_sum_title.add_run("ИТОГОВЫЙ РАСЧЕТ РЕПАРАЦИИ БИОМАССЫ (K_rec = NDVI_post / NDVI_init × 100%):")
    r_sum.bold = True
    r_sum.font.size = Pt(7.5)
    r_sum.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
    
    t_summary = doc.add_table(rows=3, cols=4)
    t_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_summary, color="0f766e", sz="4")
    set_cell_background(t_summary.rows[0].cells[0], "f1f5f9")
    set_cell_background(t_summary.rows[0].cells[1], "f1f5f9")
    set_cell_background(t_summary.rows[0].cells[2], "f1f5f9")
    set_cell_background(t_summary.rows[0].cells[3], "f1f5f9")
    
    headers_sum = ["Когорта (Кассета)", "Финальный NDVI / PLA", "Коэффициент репарации K_rec (%)", "Итог: спасено биомассы / некроз"]
    for j, h in enumerate(headers_sum):
        c = t_summary.rows[0].cells[j]
        c.paragraphs[0].add_run(h).bold = True
        c.paragraphs[0].runs[0].font.size = Pt(7.0)
        set_cell_margins(c, top=10, bottom=10, left=20, right=20)
        
    t_summary.rows[1].cells[0].paragraphs[0].add_run("К3 (Превентивная регидратация):").font.size = Pt(7.0)
    t_summary.rows[1].cells[1].paragraphs[0].add_run("NDVI = ____  PLA = ____ см²").font.size = Pt(7.0)
    t_summary.rows[1].cells[2].paragraphs[0].add_run("K_rec = ________ %").font.size = Pt(7.0)
    t_summary.rows[1].cells[3].paragraphs[0].add_run("Тургор за 2 ч, сохранено: ___ % (некроз 0%)").font.size = Pt(7.0)
    
    t_summary.rows[2].cells[0].paragraphs[0].add_run("К4 (Традиционный визуальный контроль):").font.size = Pt(7.0)
    t_summary.rows[2].cells[1].paragraphs[0].add_run("NDVI = ____  PLA = ____ см²").font.size = Pt(7.0)
    t_summary.rows[2].cells[2].paragraphs[0].add_run("K_rec = ________ %").font.size = Pt(7.0)
    t_summary.rows[2].cells[3].paragraphs[0].add_run("Краевой некроз, потеряно: ___ %").font.size = Pt(7.0)
    
    for row in t_summary.rows[1:]:
        for cell in row.cells:
            set_cell_margins(cell, top=10, bottom=10, left=20, right=20)
            
    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(3)
    p_sign.paragraph_format.space_after = Pt(0)
    r_sig = p_sign.add_run("Подпись исследователя (Алиса Ковалева): ______________________   Подпись научного руководителя: ______________________")
    r_sig.font.size = Pt(7.5)
    
    doc.save(output_path)
    print(f"[OK] Generated strictly 3-page DOCX: {output_path}")

def make_html_3pages():
    p0_days = [
        ("29.09 (Д0)", "К1–К5 (все)", "Вода 15–20 мл"),
        ("30.09 (Д1)", "К1–К5 (все)", "Пульверизатор 5–10мл"),
        ("01.10 (Д2)", "К1–К5 (все)", "Пульверизатор 5–10мл"),
        ("02.10 (Д3)", "К1–К5 (все)", "Вода 20 мл"),
        ("03.10 (Д4)", "К1–К5 + Ст0", "Вода 20 мл"),
        ("04.10 (Д5)", "К1–К5 + Ст0", "Вода 20 мл"),
        ("05.10 (Д6)", "К1–К5 + Ст0", "Вода 20 мл"),
        ("06.10 (Д7)", "К1–К5 (все)", "Вода 20 мл"),
        ("07.10 (Д8)", "К1–К5 (все)", "Вода 20 мл"),
        ("08.10 (Д9)", "К1–К5 (все)", "Вода 20 мл")
    ]
    
    cohorts_template = [
        ("№1: 🌱 Контроль (Оптимум)", "Вода 20мл", "color: #047857; font-weight: bold;"),
        ("№2: 🧂 Засоление (150 мМ)", "NaCl 20мл", "color: #b45309; font-weight: bold;"),
        ("№3: 🔬 Превентивная регидратация", "Вода «__» мл", "color: #0284c7; font-weight: bold;"),
        ("№4: 👁️ Традиционный визуальный контроль", "Вода «__» мл", "color: #6d28d9; font-weight: bold;"),
        ("№5: ⚠️ Терминал", "Полива нет", "color: #b91c1c; font-weight: bold;")
    ]

    def render_html_table(days_list):
        rows = []
        for d_lbl in days_list:
            for c_name, c_pol, c_style in cohorts_template:
                pol_color = "#b91c1c" if "нет" in c_pol else ("#0284c7" if "«__»" in c_pol else "#0f766e")
                rows.append(f"""
                <tr>
                    <td style="font-weight: bold; font-size: 7pt;">{d_lbl}</td>
                    <td style="{c_style} text-align: left; padding-left: 3px; font-size: 7.2pt;">{c_name}</td>
                    <td style="color: {pol_color}; font-weight: bold; font-size: 7pt;">{c_pol}</td>
                    <td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td>
                    <td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td>
                    <td>&nbsp;</td><td>&nbsp;</td>
                </tr>
                """)
        return "".join(rows)

    html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="UTF-8">
<title>Рабочий дневник исследователя (3 страницы) — Ковалева Алиса</title>
<style>
    @page {{
        size: A4 landscape;
        margin: 6mm 8mm 6mm 8mm;
    }}
    body {{
        font-family: Arial, sans-serif;
        font-size: 7.5pt;
        color: #0f172a;
        line-height: 1.1;
        background: #fff;
        margin: 0;
        padding: 0;
    }}
    .page-break {{
        page-break-before: always;
    }}
    h1 {{
        color: #008276;
        font-size: 10.5pt;
        margin: 0 0 2px 0;
        text-align: center;
        text-transform: uppercase;
    }}
    .subtitle {{
        text-align: center;
        font-size: 7.2pt;
        color: #475569;
        margin-bottom: 3px;
        font-style: italic;
    }}
    .guide-box {{
        width: 100%;
        border-collapse: collapse;
        margin-bottom: 3px;
        font-size: 6.8pt;
    }}
    .guide-box td {{
        padding: 3px 5px;
        vertical-align: top;
        border: 1px solid #0f766e;
    }}
    .guide-left {{
        background: #f0fdf4;
        width: 50%;
    }}
    .guide-right {{
        background: #fef2f2;
        width: 50%;
    }}
    table.data-table {{
        width: 100%;
        border-collapse: collapse;
        font-size: 6.8pt;
        margin-bottom: 2px;
    }}
    table.data-table th {{
        background: #e2e8f0;
        color: #0f766e;
        border: 1px solid #94a3b8;
        padding: 2px 1px;
        text-align: center;
        font-weight: bold;
    }}
    table.data-table td {{
        border: 1px solid #94a3b8;
        padding: 2px 1px;
        text-align: center;
        height: 13px;
    }}
    .h2-title {{
        font-size: 8.2pt;
        color: #0f766e;
        margin: 2px 0 2px 0;
        font-weight: bold;
        text-transform: uppercase;
    }}
</style>
</head>
<body>

<!-- СТРАНИЦА 1 -->
<h1>📋 РАБОЧИЙ ДНЕВНИК ИССЛЕДОВАТЕЛЯ — КОМПЛЕКС ИНДИКАЦИИ СТРЕССА РАСТЕНИЙ</h1>
<div class="subtitle">
    Исследователь: <b>Ковалева Алиса</b> (10 кл., ГБОУ СОШ №282 СПб) | Науч. рук.: Ковалев И. В. | Посев: <b>29.09.2026</b> | Старт стресс-опыта: <b>09.10.2026</b>
</div>

<table class="guide-box">
    <tr>
        <td class="guide-left">
            <b style="color: #047857;">🌱 ФАЗА 0: ПРОРАЩИВАНИЕ (29.09 — 08.10) — ЧТО ДЕЛАТЬ И СМОТРЕТЬ:</b><br>
            • <b>29.09 (Д0)</b>: Посев проростков с корешками 1–2 см. Полив отстоянной водой (15 мл). Кассета №2 СРАЗУ в отдельный лоток!<br>
            • <b>30.09 (Д1)</b>: Корни без листьев не транспирируют! Не заливать болото: только увлажнить пульверизатором (5–10 мл).<br>
            • <b>01.10 (Д2)</b>: СМОТРЕТЬ ПЕТЕЛЬКИ! При появлении зеленых всходов — <b>СРАЗУ СНЯТЬ ПЛЕНКУ</b>, чтобы не загнили!<br>
            • <b>02.10 (Д3)</b>: Раскрытие семядолей. Включить подсветку (14 ч/день). Полив по 20 мл водой.<br>
            • <b>03.10 (Д4)</b>: Посылка! Разворачивание 1-го настоящего листа. Монтаж светодиодов и АЦП.<br>
            • <b>04–05.10 (Д5–Д6):</b> Полив по 20 мл. На Стенде №0 настраиваем NoIR, белый эталон и OCR.<br>
            • <b>06–07.10 (Д7–Д8):</b> Полив по 20 мл. Рост 2-го настоящего листа, выравнивание проростков.<br>
            • <b>08.10 (Д9):</b> Вечер перед стартом: полив водой (20 мл), замер фона NDVI_init и массы M0.
        </td>
        <td class="guide-right">
            <b style="color: #b91c1c;">⚡ ФАЗА 1: СТРЕСС-ОПЫТ (09.10 — 16.10, замер в 18:00 на пике стресса):</b><br>
            • <b>Когорта №1 (🌱 Контроль):</b> Полив чистой водой 20 мл ЕЖЕДНЕВНО. Смотреть: тургор 100%, рост биомассы, лист прохладный (ΔT < 0).<br>
            • <b>Когорта №2 (🧂 Засоление 150 мМ):</b> Полив раствором соли 150 мМ ЕЖЕДНЕВНО (20 мл, 8.76 г NaCl/л). <b>НЕ ВЫНИМАТЬ ИЗ СВОЕГО ЛОТКА!</b> Смотреть: почва сырая, но лист горячий (ΔT > 0).<br>
            • <b>Когорта №3 (🔬 Превентивная регидратация):</b> С 09.10 <b>НЕ ПОЛИВАТЬ!</b> Ждать алерта (ΔT > +0.8°C, спад NDVI > 10%). При алерте — <b>СРОЧНО РЕАНИМАЦИОННЫЙ ПОЛИВ ВОДОЙ 25 мл ДО ПОЯВЛЕНИЯ УВЯДАНИЯ!</b><br>
            • <b>Когорта №4 (👁️ Традиционный визуальный контроль):</b> С 09.10 <b>НЕ ПОЛИВАТЬ!</b> Алерт игнорировать. Полив водой (25 мл) <b>ТОЛЬКО КОГДА ЛИСТЬЯ ВИДИМО ПОВИСНУТ</b> (угол > 30°). Смотреть: краевой некроз.<br>
            • <b>Когорта №5 (⚠️ Терминал):</b> <b>ВООБЩЕ НЕ ПОЛИВАТЬ ДО КОНЦА!</b> Фиксация точки гибели ткани.<br>
            <i>*Совет по регидратации (К3 и К4):</i> сухой торф поливать медленно шприцем под стебли, чтобы вода впиталась в ком!
        </td>
    </tr>
</table>

<div class="h2-title">ТАБЛИЦА РЕАЛЬНЫХ ЗАМЕРОВ: ФАЗА 0 — ПРОРАЩИВАНИЕ И ПОДГОТОВКА (ДНИ 0–9)</div>
<table class="data-table">
    <thead>
        <tr>
            <th style="width: 10%;">Дата / День</th>
            <th style="width: 11%;">Когорта (Кассета)</th>
            <th style="width: 8%;">Время</th>
            <th style="width: 14%;">Полив (мл / состав)</th>
            <th style="width: 8%;">Масса М (г)</th>
            <th style="width: 8%;">Всхожесть</th>
            <th style="width: 9%;">Tвозд / RHвозд</th>
            <th style="width: 26%;">Фенонаблюдения Алисы (ростки, петельки, листья)</th>
            <th style="width: 6%;">Подпись</th>
        </tr>
    </thead>
    <tbody>
        {''.join([f'''<tr>
            <td style="font-weight:bold;">{d}</td><td>{c}</td>
            <td>&nbsp;</td>
            <td style="color:{("#0284c7" if "Пульверизатор" in p else "#0f766e")}; font-weight:bold;">{p}</td>
            <td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td><td>&nbsp;</td>
        </tr>''' for d, c, p in p0_days])}
    </tbody>
</table>

<!-- СТРАНИЦА 2 -->
<div class="page-break"></div>
<div class="h2-title">СТРАНИЦА 2. ЖУРНАЛ СТРЕСС-ОПЫТА: ДНИ 10–13 (ВХОД В СТРЕСС, АЛЕРТ И ПРЕВЕНТИВНАЯ РЕГИДРАТАЦИЯ К3)</div>
<table class="data-table">
    <thead>
        <tr>
            <th style="width: 10%;">День / Время</th>
            <th style="width: 15%;">Когорта (Кассета)</th>
            <th style="width: 9%;">Полив (мл/чем)</th>
            <th style="width: 6%;">Масса (г)</th>
            <th style="width: 6%;">W почвы%</th>
            <th style="width: 5%;">V почвы</th>
            <th style="width: 6%;">T листа°C</th>
            <th style="width: 6%;">T возд°C</th>
            <th style="width: 5%;">ΔT (°C)</th>
            <th style="width: 5%;">RH%</th>
            <th style="width: 6%;">NDVI</th>
            <th style="width: 5%;">PLA см²</th>
            <th style="width: 12%;">Тургор / Симптомы (осмотр Алисы)</th>
            <th style="width: 4%;">Подпись</th>
        </tr>
    </thead>
    <tbody>
        {render_html_table(["День 10 09.10 (18:00)", "День 11 10.10 (18:00)", "День 12 11.10 (18:00)", "День 13 12.10 (18:00)"])}
    </tbody>
</table>

<!-- СТРАНИЦА 3 -->
<div class="page-break"></div>
<div class="h2-title">СТРАНИЦА 3. ЖУРНАЛ СТРЕСС-ОПЫТА: ДНИ 14–17 (РЕАКТИВНЫЙ ПОЛИВ К4, РЕПАРАЦИЯ, ФИНАЛЬНЫЙ УЧЕТ)</div>
<table class="data-table">
    <thead>
        <tr>
            <th style="width: 10%;">День / Время</th>
            <th style="width: 15%;">Когорта (Кассета)</th>
            <th style="width: 9%;">Полив (мл/чем)</th>
            <th style="width: 6%;">Масса (г)</th>
            <th style="width: 6%;">W почвы%</th>
            <th style="width: 5%;">V почвы</th>
            <th style="width: 6%;">T листа°C</th>
            <th style="width: 6%;">T возд°C</th>
            <th style="width: 5%;">ΔT (°C)</th>
            <th style="width: 5%;">RH%</th>
            <th style="width: 6%;">NDVI</th>
            <th style="width: 5%;">PLA см²</th>
            <th style="width: 12%;">Тургор / Симптомы (осмотр Алисы)</th>
            <th style="width: 4%;">Подпись</th>
        </tr>
    </thead>
    <tbody>
        {render_html_table(["День 14 13.10 (18:00)", "День 15 14.10 (18:00)", "День 16 15.10 (18:00)", "День 17 16.10 (18:00)"])}
    </tbody>
</table>

<div style="margin-top: 3px; border: 1px solid #0f766e; padding: 3px 6px; font-size: 7.0pt; background: #f8fafc;">
    <b>ИТОГОВЫЙ РАСЧЕТ РЕПАРАЦИИ БИОМАССЫ (K_rec = NDVI_post / NDVI_init × 100%):</b><br>
    • <b>К3 (Превентивная регидратация):</b> K_rec = ________ % | Сохранено продуктивности: ________ % | Видимый некроз: 0%<br>
    • <b>К4 (Традиционный визуальный контроль):</b> K_rec = ________ % | Безвозвратная потеря биомассы: ________ % | Краевой некроз: ______ %<br>
    <b>Подпись исследователя (Алиса Ковалева):</b> ______________________ &nbsp;&nbsp;&nbsp;&nbsp; <b>Подпись руководителя:</b> ______________________
</div>

</body>
</html>
"""
    return html

def main():
    print("=== ГЕНЕРАЦИЯ 3-СТРАНИЧНОГО РАБОЧЕГО ДНЕВНИКА С УЧЕТОМ РЕКОМЕНДАЦИЙ ===")
    os.makedirs(DOCS_DIR, exist_ok=True)
    os.makedirs(USER_DOCS, exist_ok=True)
    
    # 1. Генерируем DOCX
    docx_path = os.path.join(DOCS_DIR, "Рабочий_дневник_исследователя_Ковалева_Алиса_3стр.docx")
    user_docx_path = os.path.join(USER_DOCS, "Рабочий_дневник_исследователя_Ковалева_Алиса_3стр.docx")
    build_docx_3pages(docx_path)
    try:
        shutil.copy2(docx_path, user_docx_path)
        print(f"[OK] Copied 3-page DOCX: {user_docx_path}")
    except PermissionError:
        user_docx_v2 = os.path.join(USER_DOCS, "Рабочий_дневник_исследователя_Ковалева_Алиса_3стр_обновленный.docx")
        shutil.copy2(docx_path, user_docx_v2)
        print(f"[Notice] Основной DOCX открыт в Word. Сохранена обновленная копия: {user_docx_v2}")
    
    # 2. Генерируем HTML
    html_path = os.path.join(DOCS_DIR, "Рабочий_дневник_исследователя_Ковалева_Алиса_3стр.html")
    html_content = make_html_3pages()
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"[OK] Saved 3-page HTML: {html_path}")
    
    # 3. Генерируем PDF через Chrome
    pdf_path = os.path.join(DOCS_DIR, "Рабочий_дневник_исследователя_Ковалева_Алиса_3стр.pdf")
    user_pdf_path = os.path.join(USER_DOCS, "Рабочий_дневник_исследователя_Ковалева_Алиса_3стр.pdf")
    if os.path.exists(CHROME_PATH):
        cmd = [
            CHROME_PATH,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path
        ]
        res = subprocess.run(cmd, capture_output=True)
        if res.returncode == 0 and os.path.exists(pdf_path):
            try:
                shutil.copy2(pdf_path, user_pdf_path)
                print(f"[OK] Saved 3-page PDF: {pdf_path} ({os.path.getsize(pdf_path)/1024:.1f} KB)")
                print(f"[OK] Copied to Documents: {user_pdf_path}")
            except PermissionError:
                user_pdf_v2 = os.path.join(USER_DOCS, "Рабочий_дневник_исследователя_Ковалева_Алиса_3стр_обновленный.pdf")
                shutil.copy2(pdf_path, user_pdf_v2)
                print(f"[Notice] Основной PDF открыт в просмотрщике. Сохранена обновленная копия: {user_pdf_v2}")
            
    # Проверяем количество страниц в PDF
    if os.path.exists(user_pdf_path):
        with open(user_pdf_path, 'rb') as f:
            pdf_data = f.read()
        pages = re.findall(rb'/Type\s*/Page\b', pdf_data)
        print(f"\n>>> ПРОВЕРКА ОБЪЕМА: В PDF СТРОГО {len(pages)} СТРАНИЦЫ! <<<")

if __name__ == "__main__":
    main()
