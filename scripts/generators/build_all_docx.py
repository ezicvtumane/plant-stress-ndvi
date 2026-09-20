# -*- coding: utf-8 -*-
"""
Универсальный генератор и конвертер документов DOCX для Ковалевой Алисы.
Создает редактируемые документы Microsoft Word (.docx) для всех материалов проекта:
- Конкурсная работа «Большие вызовы»
- Научно-исследовательская работа СПбГУ
- Краткая записка для рецензирования СПбГУ
- Ответы на каверзные вопросы жюри
- Речь на очную защиту (7 минут)
- Паспорт проекта СПбГУ
- Тезисы доклада СПбГУ
- Научно-технический паспорт «Сириус»
- Пошаговый алгоритм исследования
- Аналитический отчет о проделанной работе
"""

import os
import re
import io
import base64
import shutil
import tempfile
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from bs4 import BeautifulSoup
import markdown

BASE_DIR = r"c:\Users\Администратор\Documents\Coglet\plant-stress-ndvi"
DOCS_DIR = os.path.join(BASE_DIR, "docs")
USER_DOCS = r"C:\Users\Администратор\Documents"

def set_cell_background(cell, hex_color):
    """Устанавливает цвет фона ячейки таблицы."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Устанавливает внутренние отступы ячейки."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)

def set_table_borders(table, color="cbd5e1", sz="4", val="single"):
    """Устанавливает тонкие аккуратные рамки для таблицы."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(borders)

def create_base_docx(title=""):
    """Создает новый документ Word с академическими полями и стилями."""
    doc = docx.Document()
    
    # Поля страницы: Верх 20мм, Низ 20мм, Лево 25мм, Право 15мм
    for section in doc.sections:
        section.top_margin = Inches(0.79)     # 20 мм
        section.bottom_margin = Inches(0.79)  # 20 мм
        section.left_margin = Inches(0.98)    # 25 мм
        section.right_margin = Inches(0.59)   # 15 мм
        
    # Базовый стиль
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11.5)
    font.color.rgb = RGBColor(0x1e, 0x29, 0x3b) # тёмно-серый
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)
    
    return doc

def add_callout(doc, text, alert_type="note"):
    """Добавляет блок-врезку (цитату) с боковой полосой."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    
    border_color = "00a499" # бирюзовый
    if "tip" in alert_type.lower() or "совет" in alert_type.lower():
        border_color = "059669" # зеленый
    elif "warn" in alert_type.lower() or "предупр" in alert_type.lower():
        border_color = "d97706" # янтарный
        
    p_pr = p._p.get_or_add_pPr()
    p_bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="12" w:color="{border_color}"/></w:pBdr>')
    p_pr.append(p_bdr)
    
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="f0fdfa"/>')
    p_pr.append(shd)
    
    run = p.add_run(text)
    run.font.italic = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)

def html_to_docx(html_content, output_docx_path, doc_title=""):
    """Конвертирует структурированный HTML в красивый документ DOCX."""
    doc = create_base_docx(doc_title)
    soup = BeautifulSoup(html_content, 'html.parser')
    
    # Находим основной контейнер или body
    body = soup.find('body') or soup
    
    # Пропускаем стили, скрипты, навигацию
    for s in body(['style', 'script', 'nav']):
        s.decompose()
        
    # Временная директория для изображений
    temp_dir = tempfile.mkdtemp()
    
    try:
        def process_element(elem):
            if elem.name in ['h1', 'h2', 'h3', 'h4']:
                text = elem.get_text().strip()
                if not text:
                    return
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                
                run = p.add_run(text)
                run.bold = True
                
                if elem.name == 'h1':
                    p.paragraph_format.space_before = Pt(14)
                    p.paragraph_format.space_after = Pt(8)
                    run.font.size = Pt(16)
                    run.font.color.rgb = RGBColor(0x00, 0x82, 0x76) # Sirius dark teal
                elif elem.name == 'h2':
                    p.paragraph_format.space_before = Pt(12)
                    p.paragraph_format.space_after = Pt(6)
                    run.font.size = Pt(13.5)
                    run.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
                elif elem.name == 'h3':
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(4)
                    run.font.size = Pt(12)
                    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                elif elem.name == 'h4':
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(2)
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
                    
            elif elem.name == 'p':
                text = elem.get_text().strip()
                if not text and not elem.find('img'):
                    return
                
                # Проверка на выравнивание и классы
                p = doc.add_paragraph()
                classes = elem.get('class', [])
                if any('center' in c.lower() for c in classes) or elem.get('align') == 'center':
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif any('right' in c.lower() for c in classes):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    
                p.paragraph_format.space_after = Pt(4)
                
                # Обрабатываем дочерние инлайн-элементы
                for child in elem.children:
                    if isinstance(child, str):
                        p.add_run(child)
                    elif child.name in ['b', 'strong']:
                        r = p.add_run(child.get_text())
                        r.bold = True
                    elif child.name in ['i', 'em']:
                        r = p.add_run(child.get_text())
                        r.italic = True
                    elif child.name == 'code':
                        r = p.add_run(child.get_text())
                        r.font.name = 'Consolas'
                        r.font.size = Pt(10)
                        r.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
                    elif child.name == 'a':
                        r = p.add_run(child.get_text())
                        r.font.color.rgb = RGBColor(0x02, 0x84, 0xc7)
                        r.underline = True
                    elif child.name == 'img':
                        insert_image(child, doc)
                    else:
                        p.add_run(child.get_text())
                        
            elif elem.name in ['ul', 'ol']:
                is_num = (elem.name == 'ol')
                for i, li in enumerate(elem.find_all('li', recursive=False)):
                    text = li.get_text().strip()
                    if not text:
                        continue
                    p = doc.add_paragraph(style='List Number' if is_num else 'List Bullet')
                    p.paragraph_format.space_after = Pt(2)
                    for child in li.children:
                        if isinstance(child, str):
                            p.add_run(child)
                        elif child.name in ['b', 'strong']:
                            r = p.add_run(child.get_text())
                            r.bold = True
                        elif child.name in ['i', 'em']:
                            r = p.add_run(child.get_text())
                            r.italic = True
                        else:
                            p.add_run(child.get_text())
                            
            elif elem.name == 'blockquote':
                text = elem.get_text().strip()
                if text:
                    add_callout(doc, text)
                    
            elif elem.name == 'table':
                render_table(elem, doc)
                
            elif elem.name == 'img':
                insert_image(elem, doc)
                
            elif elem.name in ['div', 'section', 'article']:
                # Рекурсивный обход блоков
                for child in elem.children:
                    if child.name:
                        process_element(child)

        def insert_image(img_tag, doc):
            src = img_tag.get('src', '')
            if not src:
                return
            try:
                img_path = None
                if src.startswith('data:image'):
                    # base64 data uri
                    header, b64_data = src.split(',', 1)
                    ext = "png"
                    if "jpeg" in header or "jpg" in header: ext = "jpg"
                    elif "svg" in header: ext = "svg"
                    
                    if ext != "svg": # Word не все SVG поддерживает через add_picture
                        temp_file = os.path.join(temp_dir, f"img_{len(os.listdir(temp_dir))}.{ext}")
                        with open(temp_file, 'wb') as f:
                            f.write(base64.b64decode(b64_data))
                        img_path = temp_file
                else:
                    # Путь к локальному файлу
                    candidates = [
                        src,
                        os.path.join(BASE_DIR, src),
                        os.path.join(DOCS_DIR, src),
                        os.path.join(BASE_DIR, src.lstrip('/'))
                    ]
                    for cand in candidates:
                        if os.path.exists(cand) and not os.path.isdir(cand) and not cand.endswith('.svg'):
                            img_path = cand
                            break
                            
                if img_path and os.path.exists(img_path):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    doc.add_picture(img_path, width=Inches(5.5))
            except Exception as e:
                print(f"  [Warn] Image insert error: {e}")

        def render_table(tbl_elem, doc):
            rows = tbl_elem.find_all('tr')
            if not rows:
                return
            
            # Определяем максимальное число колонок
            max_cols = max(len(r.find_all(['th', 'td'])) for r in rows)
            table = doc.add_table(rows=len(rows), cols=max_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table)
            
            for r_idx, row in enumerate(rows):
                cells = row.find_all(['th', 'td'])
                is_header = bool(row.find('th')) or r_idx == 0
                bg_color = "f1f5f9" if is_header else ("ffffff" if r_idx % 2 == 1 else "f8fafc")
                
                for c_idx, cell_elem in enumerate(cells):
                    if c_idx >= max_cols:
                        break
                    doc_cell = table.cell(r_idx, c_idx)
                    set_cell_background(doc_cell, bg_color)
                    set_cell_margins(doc_cell, top=80, bottom=80, left=120, right=120)
                    doc_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    
                    p = doc_cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.05
                    
                    text = cell_elem.get_text().strip()
                    run = p.add_run(text)
                    if is_header:
                        run.bold = True
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(0x0f, 0x76, 0x6e)
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        run.font.size = Pt(9.5)
                        
            doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Запуск обработки
        for elem in body.children:
            if elem.name:
                process_element(elem)
                
        doc.save(output_docx_path)
        print(f"  [OK] Saved DOCX: {os.path.basename(output_docx_path)}")
        return True
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

def markdown_to_docx(md_path, output_docx_path, doc_title=""):
    """Конвертирует Markdown файл в красивый DOCX через HTML парсинг."""
    with open(md_path, 'r', encoding='utf-8', errors='ignore') as f:
        md_text = f.read()
    
    html = markdown.markdown(md_text, extensions=['tables', 'fenced_code'])
    return html_to_docx(html, output_docx_path, doc_title)

# ==============================================================================
# ОСНОВНОЙ ПАКЕТНЫЙ ЗАПУСК
# ==============================================================================

DOCUMENTS_MAP = [
    {
        "name": "1. Конкурсная работа «Большие вызовы»",
        "src": os.path.join(DOCS_DIR, "Конкурсная_работа_Большие_Вызовы_Ковалева_Алиса.html"),
        "type": "html",
        "out": "Конкурсная_работа_Большие_Вызовы_Ковалева_Алиса.docx"
    },
    {
        "name": "2. Научно-исследовательская работа СПбГУ",
        "src": os.path.join(DOCS_DIR, "Научно_исследовательская_работа_СПбГУ_Ковалева_Алиса.html"),
        "type": "html",
        "out": "Научно_исследовательская_работа_СПбГУ_Ковалева_Алиса.docx"
    },
    {
        "name": "3. Краткая записка для рецензирования СПбГУ",
        "src": os.path.join(DOCS_DIR, "Краткая_записка_для_рецензирования_СПбГУ_Ковалева_Алиса.html"),
        "type": "html",
        "out": "Краткая_записка_для_рецензирования_СПбГУ_Ковалева_Алиса.docx"
    },
    {
        "name": "4. Ответы на каверзные вопросы жюри (Шпаргалка)",
        "src": os.path.join(DOCS_DIR, "ОТВЕТЫ_НА_КАВЕРЗНЫЕ_ВОПРОСЫ_ЖЮРИ.md"),
        "type": "md",
        "out": "Ответы_на_каверзные_вопросы_жюри_Ковалева_Алиса.docx"
    },
    {
        "name": "5. Сценарий речи на очную защиту (7 минут)",
        "src": os.path.join(DOCS_DIR, "РЕЧЬ_НА_ЗАЩИТУ_7_МИНУТ.md"),
        "type": "md",
        "out": "Речь_для_защиты_Большие_Вызовы_Ковалева_Алиса.docx"
    },
    {
        "name": "6. Паспорт проекта СПбГУ",
        "src": os.path.join(DOCS_DIR, "Паспорт_проекта_СПбГУ_Ковалева_Алиса.md"),
        "type": "md",
        "out": "Паспорт_проекта_СПбГУ_Ковалева_Алиса.docx"
    },
    {
        "name": "7. Тезисы доклада для сборника СПбГУ",
        "src": os.path.join(DOCS_DIR, "Тезисы_доклада_СПбГУ_Ковалева_Алиса.md"),
        "type": "md",
        "out": "Тезисы_доклада_СПбГУ_Ковалева_Алиса.docx"
    },
    {
        "name": "8. Научно-технический паспорт проекта «Сириус»",
        "src": os.path.join(DOCS_DIR, "SCIENTIFIC_PASSPORT.md"),
        "type": "md",
        "out": "Научно_технический_паспорт_проекта_Сириус.docx"
    },
    {
        "name": "9. Пошаговый алгоритм выполнения исследования",
        "src": os.path.join(DOCS_DIR, "STEP_BY_STEP_ALGORITHM.md"),
        "type": "md",
        "out": "Пошаговый_алгоритм_исследования.docx"
    },
    {
        "name": "10. Аналитический отчет о проделанной работе",
        "src": os.path.join(DOCS_DIR, "analysis_report.html"),
        "type": "html",
        "out": "Анализ_проделанной_работы_Комплекс_NDVI.docx"
    }
]

def main():
    print("=== ГЕНЕРАЦИЯ РЕДАКТИРУЕМЫХ ДОКУМЕНТОВ MICROSOFT WORD (.DOCX) ===")
    os.makedirs(DOCS_DIR, exist_ok=True)
    os.makedirs(USER_DOCS, exist_ok=True)
    
    success_count = 0
    generated_files = []
    
    for item in DOCUMENTS_MAP:
        print(f"\nКонвертация: {item['name']}...")
        src_path = item['src']
        out_name = item['out']
        out_docs_path = os.path.join(DOCS_DIR, out_name)
        out_user_path = os.path.join(USER_DOCS, out_name)
        
        if not os.path.exists(src_path):
            print(f"  [Skip] Файл источника не найден: {src_path}")
            continue
            
        try:
            if item['type'] == 'html':
                with open(src_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read()
                html_to_docx(content, out_docs_path, item['name'])
            elif item['type'] == 'md':
                markdown_to_docx(src_path, out_docs_path, item['name'])
                
            # Дублируем файл в пользовательскую папку Документы
            shutil.copy2(out_docs_path, out_user_path)
            print(f"  [OK] Скопировано в Документы: {out_user_path}")
            
            success_count += 1
            generated_files.append((out_name, os.path.getsize(out_docs_path)))
        except Exception as e:
            print(f"  [Error] Сбой конвертации: {e}")
            import traceback
            traceback.print_exc()

    print(f"\n=======================================================")
    print(f"Успешно создано {success_count} из {len(DOCUMENTS_MAP)} документов Word (.docx)!")
    print(f"Файлы размещены в:\n  1. {DOCS_DIR}\n  2. {USER_DOCS}")
    print("=======================================================")

if __name__ == "__main__":
    main()
